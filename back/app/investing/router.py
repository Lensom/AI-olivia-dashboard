from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html import unescape
from typing import Annotated
import json
import re
import time
import uuid
from urllib.parse import quote
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from psycopg import Connection
from psycopg.types.json import Json

from app.auth.router import get_current_user, get_conn
from app.auth.schemas import UserResponse
from app.investing.schemas import (
    InvestingActivitiesResponse,
    InvestingActivitiesUpdateRequest,
    InvestingActivityYearSummary,
    InvestingCapitalEntriesResponse,
    InvestingCapitalEntriesUpdateRequest,
    InvestingCapitalEntry,
    InvestingHoldingImportRow,
    InvestingHoldingsImportResponse,
    InvestingHoldingsResponse,
    InvestingHoldingsUpdateRequest,
    InvestingMetricsResponse,
    InvestingMetricsUpdateRequest,
    InvestingQuotesResponse,
    InvestingProfilesResponse,
    InvestingTickerProfile,
    InvestingNotesResponse,
    InvestingNotesUpdateRequest,
    InvestingPickingResponse,
    InvestingPickingRow,
    InvestingPickingUpdateRequest,
    InvestingRefillRow,
    InvestingRefillsResponse,
    InvestingRefillsUpdateRequest,
    CryptoHolding,
    CryptoPurchaseRow,
    CryptoRefillRow,
    CryptoRules,
    InvestingCryptoResponse,
    InvestingCryptoUpdateRequest,
)


router = APIRouter(prefix="/investing", tags=["investing"])
_PRICE_CACHE_TTL_SECONDS = 120
_FUNDAMENTALS_CACHE_TTL_SECONDS = 60 * 60
_price_cache: dict[str, tuple[float, float]] = {}
_fundamentals_cache: dict[str, tuple[float | None, float | None, float]] = {}
_fx_cache: dict[str, tuple[float, float]] = {}
DEFAULT_INVESTING_NOTES = """* Distribution of money ~ 15% of shares
* Main goal - general money to ETFs
* Operating expenses are no more than 3,5%
* Minimum target - 6% per year

When to buy stocks?
- During earnings season, look for a company that is undervalued and buy
- Sell bonds, spend cash

When there is nothing to buy:
- Buy value ETFs, bonds, collect cash

Important:
* 20% of obligations - after 100k in portfolio
* After 100k portfolio, increase ETF from 40% to 50% (15/15/15/5)
"""


def _parse_num(value: str) -> float:
    raw = str(value).strip().replace(" ", "")
    if "," in raw and "." not in raw:
        # Common EU format: 1090,66 -> 1090.66
        raw = raw.replace(",", ".")
    elif "," in raw and "." in raw:
        # Assume commas are thousands separators: 1,090.66 -> 1090.66
        raw = raw.replace(",", "")
    cleaned = "".join(ch for ch in raw if ch.isdigit() or ch in ".-")
    try:
        return float(cleaned) if cleaned else 0.0
    except Exception:
        return 0.0


def _fmt_pct(value: float) -> str:
    return f"{value:.2f}%"


def _parse_date(value: str) -> datetime:
    raw = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt)
        except Exception:
            pass
    return datetime.min


def _enrich_entries(entries: list[dict]) -> list[InvestingCapitalEntry]:
    enriched: list[InvestingCapitalEntry] = []
    sorted_entries = sorted(entries, key=lambda e: _parse_date(str(e.get("date", ""))))
    prev_total = 0.0
    for raw in sorted_entries:
        deposit = _parse_num(str(raw.get("deposit", "0")))
        total_value = _parse_num(str(raw.get("total_value", "0")))

        # "%" column: how total value changed between start/end of year.
        # (start ~ previous row total, end ~ current row total).
        roi = ((total_value - prev_total) / prev_total * 100.0) if prev_total > 0 else 0.0

        # "1Y" column: portfolio performance for the year accounting for top-ups.
        # First row uses only current deposit as base.
        # Next rows use previous total + current deposit.
        base = deposit if prev_total <= 0 else (prev_total + deposit)
        one_year = ((total_value - base) / base * 100.0) if base > 0 else 0.0

        prev_total = total_value
        enriched.append(
            InvestingCapitalEntry(
                id=str(raw.get("id", "")),
                date=str(raw.get("date", "")),
                deposit=str(raw.get("deposit", "")),
                total_value=str(raw.get("total_value", "")),
                roi_percent=_fmt_pct(roi),
                one_year_percent=_fmt_pct(one_year),
            )
        )
    return enriched


def _date_to_iso(value: str) -> str:
    dt = _parse_date(value)
    if dt != datetime.min:
        return dt.strftime("%Y-%m-%d")
    return datetime.utcnow().strftime("%Y-%m-%d")


def _fetch_fx_rate_to_usd(currency: str, trade_date: str) -> float | None:
    cur = str(currency or "").strip().upper()
    if not cur:
        return None
    if cur == "USD":
        return 1.0

    iso_date = _date_to_iso(trade_date)
    cache_key = f"{cur}:{iso_date}"
    now = time.time()
    cached = _fx_cache.get(cache_key)
    if cached and now - cached[1] < 60 * 60 * 6:
        return cached[0]

    urls = [
        f"https://api.frankfurter.app/{iso_date}?from={quote(cur)}&to=USD",
        f"https://api.frankfurter.app/latest?from={quote(cur)}&to=USD",
    ]
    for url in urls:
        req = Request(url, headers={"User-Agent": "Mozilla/5.0", "Accept": "application/json"})
        try:
            with urlopen(req, timeout=6) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError):
            continue
        rates = payload.get("rates", {})
        rate = rates.get("USD")
        if isinstance(rate, (int, float)) and float(rate) > 0:
            value = float(rate)
            _fx_cache[cache_key] = (value, now)
            return value
    return None


def _enrich_activities_with_usd(rows: list[dict]) -> list[dict]:
    enriched: list[dict] = []
    for row in rows:
        raw_type = str(row.get("type", "")).strip().lower()
        kind = "sell" if raw_type in {"sell", "sale", "sold"} else "buy"
        count = int(_parse_num(str(row.get("count", "0"))))
        legacy_buy_price = str(row.get("buyPrice", "0"))
        legacy_sell_price = str(row.get("sellPrice", "0"))
        price_value = _parse_num(str(row.get("price", "")))
        if price_value <= 0:
            fallback_price = legacy_sell_price if kind == "sell" else legacy_buy_price
            price_value = _parse_num(fallback_price)

        amount_value = _parse_num(str(row.get("amount", "")))
        amount = amount_value if amount_value > 0 else (count * price_value)

        gains_value = _parse_num(str(row.get("gains_losses", "")))
        if gains_value == 0:
            gains_value = _parse_num(str(row.get("result", "0")))
        gains_losses = gains_value
        commission = _parse_num(str(row.get("commission", "0")))
        currency = str(row.get("currency", "USD")).strip().upper() or "USD"
        trade_date = str(row.get("date", "")).strip()
        fx = _fetch_fx_rate_to_usd(currency, trade_date) or (1.0 if currency == "USD" else 0.0)
        amount_usd = amount * fx
        gains_losses_usd = gains_losses * fx
        commission_usd = commission * fx
        enriched.append(
            {
                **row,
                "type": kind,
                "count": count,
                "price": f"{price_value:.6f}".rstrip("0").rstrip(".") if price_value else "0",
                "amount": f"{amount:.6f}".rstrip("0").rstrip(".") if amount else "0",
                "gains_losses": f"{gains_losses:.6f}".rstrip("0").rstrip(".") if gains_losses else "0",
                "currency": currency,
                "amount_usd": f"{amount_usd:.2f}",
                "gains_losses_usd": f"{gains_losses_usd:.2f}",
                "commission_usd": f"{commission_usd:.2f}",
                "fx_rate_to_usd": f"{fx:.6f}",
            }
        )
    return enriched


def _calc_activity_yearly_summary(rows: list[dict]) -> list[InvestingActivityYearSummary]:
    by_year: dict[str, dict[str, float]] = {}
    for row in rows:
        date = str(row.get("date", "")).strip()
        dt = _parse_date(date)
        year = str(dt.year) if dt != datetime.min else (date[-4:] if len(date) >= 4 else "Unknown")
        if year not in by_year:
            by_year[year] = {
                "purchases_amount": 0.0,
                "purchases_commission": 0.0,
                "sales_amount": 0.0,
                "sales_pnl": 0.0,
                "sales_commission": 0.0,
            }
        amount = _parse_num(str(row.get("amount_usd", row.get("amount", "0"))))
        gains_losses = _parse_num(str(row.get("gains_losses_usd", row.get("gains_losses", "0"))))
        commission = _parse_num(str(row.get("commission_usd", row.get("commission", "0"))))
        kind = str(row.get("type", "")).strip().lower()
        if kind == "sell":
            by_year[year]["sales_amount"] += amount
            by_year[year]["sales_pnl"] += gains_losses
            by_year[year]["sales_commission"] += commission
        else:
            by_year[year]["purchases_amount"] += amount
            by_year[year]["purchases_commission"] += commission

    summary: list[InvestingActivityYearSummary] = []
    for year in sorted(by_year.keys()):
        agg = by_year[year]
        net = (
            agg["sales_amount"]
            - agg["purchases_amount"]
            - agg["sales_commission"]
            - agg["purchases_commission"]
            + agg["sales_pnl"]
        )
        summary.append(
            InvestingActivityYearSummary(
                year=year,
                purchases_amount=f"{agg['purchases_amount']:.2f}",
                purchases_commission=f"{agg['purchases_commission']:.2f}",
                sales_amount=f"{agg['sales_amount']:.2f}",
                sales_pnl=f"{agg['sales_pnl']:.2f}",
                sales_commission=f"{agg['sales_commission']:.2f}",
                net_cash_flow=f"{net:.2f}",
            )
        )
    return summary


def _fetch_yahoo_prices(tickers: list[str]) -> dict[str, float]:
    rows = _fetch_yahoo_quote_rows(tickers)
    prices: dict[str, float] = {}
    for symbol, item in rows.items():
        price = item.get("regularMarketPrice")
        if isinstance(price, (int, float)):
            prices[symbol] = float(price)
    return prices


def _fetch_stooq_price(ticker: str) -> float | None:
    symbol = ticker.lower().strip()
    if not symbol:
        return None
    url = f"https://stooq.com/q/l/?s={quote(symbol)}.us&i=d"
    req = Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0",
            "Accept": "text/csv,text/plain,*/*",
        },
    )
    try:
        with urlopen(req, timeout=6) as response:
            raw = response.read().decode("utf-8", "ignore").strip()
    except (HTTPError, URLError, TimeoutError):
        return None
    if not raw:
        return None
    first = raw.splitlines()[0]
    parts = [p.strip() for p in first.split(",")]
    # Expected format: SYMBOL,DATE,TIME,OPEN,HIGH,LOW,CLOSE,VOLUME
    if len(parts) < 7:
        return None
    close = _parse_num(parts[6])
    return close if close > 0 else None


def _fetch_market_prices(tickers: list[str]) -> dict[str, float]:
    now = time.time()
    result: dict[str, float] = {}
    missing: list[str] = []
    for ticker in tickers:
        cached = _price_cache.get(ticker)
        if cached and now - cached[1] < _PRICE_CACHE_TTL_SECONDS:
            result[ticker] = cached[0]
        else:
            missing.append(ticker)

    if missing:
        yahoo_prices = _fetch_yahoo_prices(missing)
        need_stooq: list[str] = []
        for ticker in missing:
            price = yahoo_prices.get(ticker)
            if isinstance(price, (int, float)) and float(price) > 0:
                value = float(price)
                result[ticker] = value
                _price_cache[ticker] = (value, now)
            else:
                need_stooq.append(ticker)
        # Parallel so an unreachable Stooq costs one timeout, not one per ticker.
        with ThreadPoolExecutor(max_workers=_QUOTE_WORKERS) as pool:
            stooq_prices = list(pool.map(_fetch_stooq_price, need_stooq))
        for ticker, stooq_price in zip(need_stooq, stooq_prices):
            if isinstance(stooq_price, (int, float)) and float(stooq_price) > 0:
                value = float(stooq_price)
                result[ticker] = value
                _price_cache[ticker] = (value, now)
    return result


_QUOTE_WORKERS = 10


def _fetch_yahoo_chart_quote(ticker: str) -> dict | None:
    """One ticker via Yahoo v8 chart (the v7 bulk quote API now requires a crumb and returns nothing)."""
    url = f"https://query1.finance.yahoo.com/v8/finance/chart/{quote(ticker)}?range=1d&interval=1d"
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urlopen(req, timeout=8) as response:
            payload = json.loads(response.read().decode("utf-8"))
        meta = payload["chart"]["result"][0]["meta"]
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError, KeyError, IndexError, TypeError):
        return None
    price = meta.get("regularMarketPrice")
    if not isinstance(price, (int, float)):
        return None
    row: dict = {"regularMarketPrice": float(price)}
    prev = meta.get("chartPreviousClose")
    if isinstance(prev, (int, float)) and prev > 0:
        row["regularMarketChangePercent"] = (float(price) / float(prev) - 1) * 100
    return row


def _fetch_yahoo_quote_rows(tickers: list[str]) -> dict[str, dict]:
    """Ticker (uppercase) → {regularMarketPrice, regularMarketChangePercent?}, fetched in parallel."""
    if not tickers:
        return {}
    with ThreadPoolExecutor(max_workers=_QUOTE_WORKERS) as pool:
        results = list(pool.map(_fetch_yahoo_chart_quote, tickers))
    return {t.upper().strip(): row for t, row in zip(tickers, results) if row}


# Base symbol (BTC, ETH) → CoinGecko id for /simple/price (fills gaps when Yahoo/Stooq miss crypto).
_COINGECKO_IDS: dict[str, str] = {
    "BTC": "bitcoin",
    "ETH": "ethereum",
    "SOL": "solana",
    "BNB": "binancecoin",
    "XRP": "ripple",
    "ADA": "cardano",
    "DOGE": "dogecoin",
    "DOT": "polkadot",
    "MATIC": "matic-network",
    "AVAX": "avalanche-2",
    "LINK": "chainlink",
    "ATOM": "cosmos",
    "LTC": "litecoin",
    "TRX": "tron",
    "SHIB": "shiba-inu",
    "UNI": "uniswap",
    "ARB": "arbitrum",
    "OP": "optimism",
    "APT": "aptos",
    "NEAR": "near",
}


def _ticker_base_crypto(t: str) -> str:
    u = t.upper().strip()
    if u.endswith("-USD"):
        return u[:-4]
    return u


def _fetch_coingecko_crypto_quotes(tickers: list[str]) -> tuple[dict[str, float], dict[str, float]]:
    """Prices (USD) and ~24h change % from CoinGecko. Keys match requested tickers (e.g. BTC-USD)."""
    prices: dict[str, float] = {}
    day_pct: dict[str, float] = {}
    id_to_tickers: dict[str, list[str]] = {}
    for t in tickers:
        base = _ticker_base_crypto(t)
        cg_id = _COINGECKO_IDS.get(base)
        if cg_id:
            id_to_tickers.setdefault(cg_id, []).append(t)
    if not id_to_tickers:
        return prices, day_pct
    ids_param = ",".join(sorted(id_to_tickers.keys()))
    url = (
        "https://api.coingecko.com/api/v3/simple/price"
        f"?ids={ids_param}&vs_currencies=usd&include_24hr_change=true"
    )
    try:
        with urlopen(url, timeout=12) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError):
        return prices, day_pct
    if not isinstance(payload, dict):
        return prices, day_pct
    for cg_id, tlist in id_to_tickers.items():
        row = payload.get(cg_id)
        if not isinstance(row, dict):
            continue
        usd = row.get("usd")
        ch = row.get("usd_24h_change")
        if isinstance(usd, (int, float)) and float(usd) > 0:
            for t in tlist:
                prices[t] = float(usd)
        if isinstance(ch, (int, float)):
            for t in tlist:
                day_pct[t] = float(ch)
    return prices, day_pct


_PICKING_STORED_KEYS = (
    "id",
    "bucket",
    "name",
    "ticker",
    "industry",
    "pe",
    "eps",
    "beta",
    "div_yield",
    "strong_buy_until",
    "may_buy_until",
    "buy_right_now",
    "price_goal_1y",
    "price_goal_5y",
    "notes",
)


def _normalize_stored_picking_row(raw: dict) -> dict:
    d = {str(k): "" if v is None else str(v) for k, v in raw.items()}
    notes = str(d.get("notes") or "").strip()
    if not notes and d.get("reports"):
        notes = str(d.get("reports") or "").strip()
    d.pop("reports", None)
    d.pop("current_price", None)
    d.pop("day_change_pct", None)
    out = {k: "" for k in _PICKING_STORED_KEYS}
    for k in _PICKING_STORED_KEYS:
        if k == "notes":
            out[k] = notes
        else:
            out[k] = str(d.get(k) or "")
    if not str(out.get("bucket") or "").strip():
        out["bucket"] = "value_blue_chips"
    return out


def _pick_row_to_storage(row: InvestingPickingRow) -> dict:
    d = row.model_dump()
    return {k: str(d.get(k) or "") for k in _PICKING_STORED_KEYS}


def _enrich_picking_rows(rows: list[dict]) -> list[dict]:
    tickers = list(
        dict.fromkeys(
            str(r.get("ticker", "")).strip().upper() for r in rows if str(r.get("ticker", "")).strip()
        )
    )[:50]
    quotes = _fetch_yahoo_quote_rows(tickers) if tickers else {}
    missing: list[str] = []
    for t in tickers:
        q = quotes.get(t)
        rm = q.get("regularMarketPrice") if q else None
        if not isinstance(rm, (int, float)) or float(rm) <= 0:
            missing.append(t)
    backup = _fetch_market_prices(missing) if missing else {}
    mapped_cg = [t for t in tickers if _COINGECKO_IDS.get(_ticker_base_crypto(t))]
    cg_prices_p, cg_day_p = _fetch_coingecko_crypto_quotes(mapped_cg) if mapped_cg else ({}, {})
    out: list[dict] = []
    for r in rows:
        row = dict(r)
        t = str(row.get("ticker", "")).strip().upper()
        q = quotes.get(t) if t else None
        price_val: float | None = None
        day_pct: float | None = None
        if q:
            rm = q.get("regularMarketPrice")
            if isinstance(rm, (int, float)) and float(rm) > 0:
                price_val = float(rm)
            ch = q.get("regularMarketChangePercent")
            if isinstance(ch, (int, float)):
                day_pct = float(ch)
        if price_val is None and t:
            b = backup.get(t)
            if isinstance(b, (int, float)) and float(b) > 0:
                price_val = float(b)
        if price_val is None and t:
            cp = cg_prices_p.get(t)
            if isinstance(cp, (int, float)) and float(cp) > 0:
                price_val = float(cp)
        if day_pct is None and t:
            cd = cg_day_p.get(t)
            if isinstance(cd, (int, float)):
                day_pct = float(cd)
        row["current_price"] = f"{price_val:.2f}" if price_val is not None else ""
        row["day_change_pct"] = f"{day_pct:.2f}%" if day_pct is not None else ""
        out.append(row)
    return out


_BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36"
)
# Failed fetches are retried sooner than successful ones are refreshed.
_FUNDAMENTALS_FAILURE_TTL_SECONDS = 5 * 60


def _finviz_field(html: str, boxover: str) -> str:
    """Text of the first <b> value right after a Finviz snapshot label (e.g. "1.08", "3.64 (0.68%)", "-")."""
    m = re.search(rf'data-boxover-html="{re.escape(boxover)}".*?<b>([^<]*)</b>', html, re.S)
    return m.group(1).strip() if m else ""


def _fetch_finviz_fundamentals(ticker: str) -> tuple[float | None, float | None]:
    """(dividend yield as a fraction, beta). None = unknown; a non-payer's yield is 0.0, not None."""
    now = time.time()
    cached = _fundamentals_cache.get(ticker)
    if cached:
        failed = cached[0] is None and cached[1] is None
        ttl = _FUNDAMENTALS_FAILURE_TTL_SECONDS if failed else _FUNDAMENTALS_CACHE_TTL_SECONDS
        if now - cached[2] < ttl:
            return cached[0], cached[1]

    url = f"https://finviz.com/quote.ashx?t={quote(ticker)}&p=d"
    # Finviz rejects bare "Mozilla/5.0"; it needs a full browser User-Agent.
    req = Request(url, headers={"User-Agent": _BROWSER_UA, "Accept": "text/html,application/xhtml+xml,*/*"})
    try:
        with urlopen(req, timeout=7) as response:
            html = response.read().decode("utf-8", "ignore")
    except (HTTPError, URLError, TimeoutError):
        _fundamentals_cache[ticker] = (None, None, now)
        return None, None

    beta_text = _finviz_field(html, "Beta")
    beta = _parse_num(beta_text) if beta_text and beta_text != "-" else None
    if isinstance(beta, float) and beta <= 0:
        beta = None

    div_yield: float | None = None
    if 'data-boxover-html="Trailing 12 Months Dividend"' in html:
        pct = re.search(r"\(([-+]?\d+(?:\.\d+)?)%\)", _finviz_field(html, "Trailing 12 Months Dividend"))
        div_yield = float(pct.group(1)) / 100.0 if pct else 0.0

    _fundamentals_cache[ticker] = (div_yield, beta, now)
    return div_yield, beta


def _calc_metrics_from_holdings(holdings: list[dict]) -> tuple[str, str, dict[str, str]]:
    tickers: list[str] = []
    shares_by_ticker: dict[str, float] = {}
    for row in holdings:
        ticker = str(row.get("ticker", "")).strip().upper()
        if not ticker:
            continue
        shares = _parse_num(str(row.get("shares", "0")))
        shares_by_ticker[ticker] = shares_by_ticker.get(ticker, 0.0) + max(0.0, shares)
        tickers.append(ticker)
    unique_tickers = list(dict.fromkeys(tickers))[:50]
    if not unique_tickers:
        return "—", "—", {}

    prices = _fetch_market_prices(unique_tickers)
    values_by_ticker: dict[str, float] = {}
    total_value = 0.0
    for ticker in unique_tickers:
        price = prices.get(ticker)
        if not isinstance(price, (int, float)):
            continue
        value = float(price) * shares_by_ticker.get(ticker, 0.0)
        if value <= 0:
            continue
        values_by_ticker[ticker] = value
        total_value += value

    if total_value <= 0:
        return "—", "—", {}

    weighted_div_yield = 0.0
    weighted_beta = 0.0
    div_weight_total = 0.0
    beta_weight_total = 0.0
    allocation_percent: dict[str, str] = {}

    tickers_with_value = list(values_by_ticker)
    with ThreadPoolExecutor(max_workers=_QUOTE_WORKERS) as pool:
        fundamentals = dict(zip(tickers_with_value, pool.map(_fetch_finviz_fundamentals, tickers_with_value)))

    for ticker, value in values_by_ticker.items():
        weight = value / total_value
        allocation_percent[ticker] = f"{(weight * 100.0):.2f}%"

        div_yield, beta = fundamentals[ticker]
        if isinstance(div_yield, (int, float)):
            weighted_div_yield += weight * float(div_yield)
            div_weight_total += weight

        if isinstance(beta, (int, float)):
            weighted_beta += weight * float(beta)
            beta_weight_total += weight

    # Average over the holdings we have data for, so a missing ticker doesn't drag the result toward 0.
    div_yield_percent = f"{(weighted_div_yield / div_weight_total * 100.0):.2f}%" if div_weight_total > 0 else "—"
    beta_value = f"{(weighted_beta / beta_weight_total):.2f}" if beta_weight_total > 0 else "—"
    return div_yield_percent, beta_value, allocation_percent


_PROFILE_CACHE_TTL_SECONDS = 24 * 60 * 60
_profile_cache: dict[str, tuple[InvestingTickerProfile | None, float]] = {}


def _fetch_finviz_profile(ticker: str) -> InvestingTickerProfile | None:
    """Company/fund name and an industry label: sector for stocks, "ETF · <category>" for funds."""
    now = time.time()
    cached = _profile_cache.get(ticker)
    if cached:
        ttl = _PROFILE_CACHE_TTL_SECONDS if cached[0] else _FUNDAMENTALS_FAILURE_TTL_SECONDS
        if now - cached[1] < ttl:
            return cached[0]

    url = f"https://finviz.com/quote.ashx?t={quote(ticker)}&p=d"
    req = Request(url, headers={"User-Agent": _BROWSER_UA, "Accept": "text/html,application/xhtml+xml,*/*"})
    try:
        with urlopen(req, timeout=7) as response:
            html = response.read().decode("utf-8", "ignore")
    except (HTTPError, URLError, TimeoutError):
        _profile_cache[ticker] = (None, now)
        return None

    # <title>MSFT - Microsoft Corp Stock Price and Quote</title>
    title = re.search(r"<title>\s*[^<]*?\s-\s(.*?)\s+Stock Price and Quote\s*</title>", html, re.S)
    name = unescape(title.group(1).strip()) if title else ""

    # Funds carry an "Asset Type" field; their sector link is not meaningful (VOO shows "Financial").
    if 'snapshot-td-label">Asset Type<' in html:
        category = re.search(r'snapshot-td-label">Category<.*?<b>(.*?)</b>', html, re.S)
        category_text = unescape(re.sub(r"<[^>]+>", "", category.group(1))).strip() if category else ""
        industry = f"ETF · {category_text}" if category_text else "ETF"
    else:
        sector = re.search(r'href="screener\?v=111&(?:amp;)?f=sec_[^"]*"[^>]*>([^<]+)</a>', html)
        industry = unescape(sector.group(1).strip()) if sector else ""

    profile = InvestingTickerProfile(name=name, industry=industry) if (name or industry) else None
    _profile_cache[ticker] = (profile, now)
    return profile


@router.get("/profiles", response_model=InvestingProfilesResponse)
def get_profiles(
    tickers: Annotated[str, Query(description="Comma-separated tickers, e.g. AAPL,VOO")] = "",
    user: Annotated[UserResponse, Depends(get_current_user)] = None,
):
    """Name and industry per ticker (Finviz), used to autofill holdings. Unknown tickers are omitted."""
    _ = user.id
    unique = list(dict.fromkeys(t.strip().upper() for t in tickers.split(",") if t.strip()))[:50]
    if not unique:
        return InvestingProfilesResponse()
    with ThreadPoolExecutor(max_workers=_QUOTE_WORKERS) as pool:
        results = list(pool.map(_fetch_finviz_profile, unique))
    return InvestingProfilesResponse(profiles={t: p for t, p in zip(unique, results) if p})


@router.get("/quotes", response_model=InvestingQuotesResponse)
def get_quotes(
    tickers: Annotated[str, Query(description="Comma-separated tickers, e.g. AAPL,MSFT")] = "",
    user: Annotated[UserResponse, Depends(get_current_user)] = None,
):
    _ = user.id
    parsed = [t.strip().upper() for t in tickers.split(",") if t.strip()]
    unique_tickers = list(dict.fromkeys(parsed))[:50]
    if not unique_tickers:
        return InvestingQuotesResponse(prices={}, day_change_percent={}, as_of=datetime.utcnow().isoformat())
    rows = _fetch_yahoo_quote_rows(unique_tickers)
    prices: dict[str, float] = {}
    day_change: dict[str, float] = {}
    missing_prices: list[str] = []
    for t in unique_tickers:
        q = rows.get(t)
        if q:
            rm = q.get("regularMarketPrice")
            if isinstance(rm, (int, float)) and float(rm) > 0:
                prices[t] = float(rm)
            ch = q.get("regularMarketChangePercent")
            if isinstance(ch, (int, float)):
                day_change[t] = float(ch)
        if t not in prices:
            missing_prices.append(t)
    if missing_prices:
        backup = _fetch_market_prices(missing_prices)
        for t, p in backup.items():
            if t not in prices and isinstance(p, (int, float)) and float(p) > 0:
                prices[t] = float(p)

    crypto_mapped = [t for t in unique_tickers if _COINGECKO_IDS.get(_ticker_base_crypto(t))]
    if crypto_mapped:
        cg_prices, cg_day = _fetch_coingecko_crypto_quotes(crypto_mapped)
        for t, p in cg_prices.items():
            if t not in prices:
                prices[t] = p
        for t, d in cg_day.items():
            if t not in day_change:
                day_change[t] = d

    return InvestingQuotesResponse(
        prices=prices,
        day_change_percent=day_change,
        as_of=datetime.utcnow().isoformat(),
    )


@router.get("/metrics", response_model=InvestingMetricsResponse)
def get_metrics(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT holdings FROM investing_holdings WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    holdings = row[0] if row and isinstance(row[0], list) else []
    try:
        div_yield_percent, beta, allocation_percent = _calc_metrics_from_holdings(holdings)
    except Exception:
        # Never fail the dashboard because of upstream quote issues or malformed holdings.
        div_yield_percent, beta, allocation_percent = "—", "—", {}
    return InvestingMetricsResponse(
        div_yield_percent=div_yield_percent,
        beta=beta,
        allocation_percent=allocation_percent,
    )


@router.put("/metrics", response_model=InvestingMetricsResponse)
def update_metrics(
    payload: InvestingMetricsUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    # Keep existing values if fields are omitted.
    current = conn.execute(
        "SELECT div_yield_percent, beta FROM investing_metrics WHERE user_id = %s",
        (user.id,),
    ).fetchone()

    div_yield = payload.div_yield_percent if payload.div_yield_percent is not None else (current[0] if current else "—")
    beta = payload.beta if payload.beta is not None else (current[1] if current else "—")

    conn.execute(
        """
        INSERT INTO investing_metrics(user_id, div_yield_percent, beta, updated_at)
        VALUES (%s, %s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET div_yield_percent = EXCLUDED.div_yield_percent,
                      beta = EXCLUDED.beta,
                      updated_at = NOW()
        """,
        (user.id, div_yield, beta),
    )
    conn.commit()
    return InvestingMetricsResponse(div_yield_percent=div_yield, beta=beta, allocation_percent={})


@router.get("/holdings", response_model=InvestingHoldingsResponse)
def get_holdings(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT holdings, updated_at FROM investing_holdings WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return InvestingHoldingsResponse(holdings=[], updated_at=None)
    return InvestingHoldingsResponse(holdings=row[0], updated_at=str(row[1]))


@router.put("/holdings", response_model=InvestingHoldingsResponse)
def update_holdings(
    payload: InvestingHoldingsUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    holdings_json = [h.model_dump() for h in payload.holdings]
    row = conn.execute(
        """
        INSERT INTO investing_holdings(user_id, holdings, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET holdings = EXCLUDED.holdings,
                      updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(holdings_json)),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    return InvestingHoldingsResponse(holdings=holdings_json, updated_at=updated_at)


# Broker export columns (case-insensitive) → import field. Share (%), Profit, Growth are derived, so ignored.
_IMPORT_COLUMNS = {
    "ticker": "ticker",
    "qty": "shares",
    "entry price": "avgBuyPrice",
    "price": "marketPrice",
    "value": "marketValue",
}


def _import_number(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    raw = re.sub(r"[\s$]", "", str(value or ""))
    if not raw:
        return None
    if "," in raw and "." not in raw:
        raw = raw.replace(",", ".")
    else:
        raw = raw.replace(",", "")
    try:
        return float(raw)
    except ValueError:
        return None


def _import_ticker(value: object) -> str:
    # "AAPL.US" → "AAPL" (Yahoo quotes use bare US tickers); other exchange suffixes are kept.
    ticker = str(value or "").strip().upper()
    return ticker[:-3] if ticker.endswith(".US") else ticker


def _read_import_table(filename: str, content: bytes) -> list[list[object]]:
    if filename.endswith(".csv"):
        import csv
        import io

        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        return [list(r) for r in csv.reader(io.StringIO(text), dialect)]

    import io

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read the spreadsheet. Save it as .xlsx.") from exc
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


@router.post("/holdings/import", response_model=InvestingHoldingsImportResponse)
async def import_holdings(
    user: Annotated[UserResponse, Depends(get_current_user)],
    file: UploadFile = File(...),
):
    """Parse a holdings spreadsheet. Nothing is persisted: the client merges rows and saves via PUT /holdings."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm", ".csv")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .csv file.")
    table = _read_import_table(filename, await file.read())

    header_idx = next(
        (i for i, r in enumerate(table) if any(str(c or "").strip().lower() == "ticker" for c in r)),
        None,
    )
    if header_idx is None:
        raise HTTPException(status_code=400, detail="No header row with a 'Ticker' column found.")
    columns: dict[str, int] = {}
    for i, cell in enumerate(table[header_idx]):
        key = _IMPORT_COLUMNS.get(str(cell or "").strip().lower())
        if key and key not in columns:
            columns[key] = i
    missing = [label for label, key in _IMPORT_COLUMNS.items() if key not in columns]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing columns: {', '.join(missing)}.")

    rows: list[InvestingHoldingImportRow] = []
    skipped: list[str] = []
    for r in table[header_idx + 1 :]:
        cell = lambda key: r[columns[key]] if columns[key] < len(r) else None  # noqa: E731
        ticker = _import_ticker(cell("ticker"))
        if not ticker:
            continue
        nums = {key: _import_number(cell(key)) for key in ("shares", "avgBuyPrice", "marketPrice", "marketValue")}
        if nums["shares"] is None or nums["avgBuyPrice"] is None:
            skipped.append(ticker)
            continue
        price = nums["marketPrice"] or 0.0
        rows.append(
            InvestingHoldingImportRow(
                ticker=ticker,
                shares=nums["shares"],
                avgBuyPrice=nums["avgBuyPrice"],
                marketPrice=price,
                marketValue=nums["marketValue"] if nums["marketValue"] is not None else price * nums["shares"],
            )
        )
    return InvestingHoldingsImportResponse(rows=rows, skipped=skipped)


@router.get("/capital-entries", response_model=InvestingCapitalEntriesResponse)
def get_capital_entries(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT entries, updated_at FROM investing_capital_entries WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return InvestingCapitalEntriesResponse(entries=[], updated_at=None)
    entries_raw = row[0] if isinstance(row[0], list) else []
    return InvestingCapitalEntriesResponse(entries=_enrich_entries(entries_raw), updated_at=str(row[1]))


@router.put("/capital-entries", response_model=InvestingCapitalEntriesResponse)
def update_capital_entries(
    payload: InvestingCapitalEntriesUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    entries_json = [e.model_dump() for e in payload.entries]
    row = conn.execute(
        """
        INSERT INTO investing_capital_entries(user_id, entries, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET entries = EXCLUDED.entries,
                      updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(entries_json)),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    return InvestingCapitalEntriesResponse(entries=_enrich_entries(entries_json), updated_at=updated_at)


@router.get("/notes", response_model=InvestingNotesResponse)
def get_notes(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT notes, updated_at FROM investing_notes WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if row:
        return InvestingNotesResponse(notes=row[0], updated_at=str(row[1]))

    inserted = conn.execute(
        """
        INSERT INTO investing_notes(user_id, notes, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET notes = investing_notes.notes
        RETURNING notes, updated_at
        """,
        (user.id, DEFAULT_INVESTING_NOTES),
    ).fetchone()
    conn.commit()
    if not inserted:
        return InvestingNotesResponse(notes=DEFAULT_INVESTING_NOTES, updated_at=None)
    return InvestingNotesResponse(notes=inserted[0], updated_at=str(inserted[1]))


@router.put("/notes", response_model=InvestingNotesResponse)
def update_notes(
    payload: InvestingNotesUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        """
        INSERT INTO investing_notes(user_id, notes, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET notes = EXCLUDED.notes,
                      updated_at = NOW()
        RETURNING notes, updated_at
        """,
        (user.id, payload.notes),
    ).fetchone()
    conn.commit()
    if not row:
        return InvestingNotesResponse(notes=payload.notes, updated_at=None)
    return InvestingNotesResponse(notes=row[0], updated_at=str(row[1]))


@router.get("/activities", response_model=InvestingActivitiesResponse)
def get_activities(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT activities, updated_at FROM investing_activities WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return InvestingActivitiesResponse(activities=[], yearly_summary=[], updated_at=None)
    raw = row[0] if isinstance(row[0], list) else []
    try:
        enriched = _enrich_activities_with_usd(raw)
        summary = _calc_activity_yearly_summary(enriched)
    except Exception:
        # Fail-safe: never break Activities page due to malformed legacy rows.
        enriched = []
        summary = []
    return InvestingActivitiesResponse(activities=enriched, yearly_summary=summary, updated_at=str(row[1]))


@router.put("/activities", response_model=InvestingActivitiesResponse)
def update_activities(
    payload: InvestingActivitiesUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    activities_json = [
        {
            k: v
            for k, v in a.model_dump().items()
            if k not in {"amount_usd", "gains_losses_usd", "commission_usd", "fx_rate_to_usd"}
        }
        for a in payload.activities
    ]
    for row in activities_json:
        raw_type = str(row.get("type", "")).strip().lower()
        row["type"] = "sell" if raw_type in {"sell", "sale", "sold"} else "buy"
    row = conn.execute(
        """
        INSERT INTO investing_activities(user_id, activities, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET activities = EXCLUDED.activities,
                      updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(activities_json)),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    enriched = _enrich_activities_with_usd(activities_json)
    summary = _calc_activity_yearly_summary(enriched)
    return InvestingActivitiesResponse(activities=enriched, yearly_summary=summary, updated_at=updated_at)


@router.get("/picking", response_model=InvestingPickingResponse)
def get_picking(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT rows, updated_at FROM investing_picking WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return InvestingPickingResponse(rows=[], updated_at=None)
    raw = row[0] if isinstance(row[0], list) else []
    base = [_normalize_stored_picking_row(r) for r in raw if isinstance(r, dict)]
    enriched = _enrich_picking_rows(base)
    return InvestingPickingResponse(rows=[InvestingPickingRow(**e) for e in enriched], updated_at=str(row[1]))


@router.put("/picking", response_model=InvestingPickingResponse)
def update_picking(
    payload: InvestingPickingUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    rows_json = [_pick_row_to_storage(r) for r in payload.rows]
    row = conn.execute(
        """
        INSERT INTO investing_picking(user_id, rows, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET rows = EXCLUDED.rows,
                      updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(rows_json)),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    enriched = _enrich_picking_rows([dict(r) for r in rows_json])
    return InvestingPickingResponse(rows=[InvestingPickingRow(**e) for e in enriched], updated_at=updated_at)


def _normalize_refill_row(raw: dict) -> dict:
    amt = raw.get("amount")
    if amt is None and raw.get("invest") is not None:
        amt = raw.get("invest")
    return {
        "id": str(raw.get("id", "")),
        "date": str(raw.get("date", "")),
        "amount": str(amt if amt is not None else ""),
        "commission": str(raw.get("commission", "")),
    }


def _sort_refill_rows(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: _parse_date(str(r.get("date", ""))))


@router.get("/refills", response_model=InvestingRefillsResponse)
def get_refills(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT rows, operating_expenses_usd, updated_at FROM investing_refills WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return InvestingRefillsResponse(rows=[], operating_expenses_usd="", updated_at=None)
    raw_rows = row[0] if isinstance(row[0], list) else []
    opx = str(row[1] if row[1] is not None else "")
    normalized = [_normalize_refill_row(r) for r in raw_rows if isinstance(r, dict)]
    sorted_rows = _sort_refill_rows(normalized)
    return InvestingRefillsResponse(
        rows=[InvestingRefillRow(**r) for r in sorted_rows],
        operating_expenses_usd=opx,
        updated_at=str(row[2]),
    )


@router.put("/refills", response_model=InvestingRefillsResponse)
def update_refills(
    payload: InvestingRefillsUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    rows_json = [_normalize_refill_row(r.model_dump()) for r in payload.rows]
    rows_json = _sort_refill_rows(rows_json)
    opx = str(payload.operating_expenses_usd or "").strip()[:64]
    row = conn.execute(
        """
        INSERT INTO investing_refills(user_id, rows, operating_expenses_usd, updated_at)
        VALUES (%s, %s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET rows = EXCLUDED.rows,
                      operating_expenses_usd = EXCLUDED.operating_expenses_usd,
                      updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(rows_json), opx),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    return InvestingRefillsResponse(
        rows=[InvestingRefillRow(**r) for r in rows_json],
        operating_expenses_usd=opx,
        updated_at=updated_at,
    )


_CRYPTO_DEFAULT_RULES: dict[str, str] = {
    "btc_percent": "70",
    "eth_percent": "30",
    "btc_note": "",
    "eth_note": "",
}


def _normalize_crypto_holding(raw: dict) -> dict:
    rid = str(raw.get("id", "")).strip()
    return {
        "id": rid if rid else str(uuid.uuid4()),
        "symbol": str(raw.get("symbol", "")).strip(),
        "name": str(raw.get("name", "")).strip(),
        "quantity": str(raw.get("quantity", "")).strip(),
    }


def _normalize_crypto_purchase(raw: dict) -> dict:
    rid = str(raw.get("id", "")).strip()
    return {
        "id": rid if rid else str(uuid.uuid4()),
        "date": str(raw.get("date", "")).strip(),
        "btc_usd": str(raw.get("btc_usd", "")).strip(),
        "btc_price": str(raw.get("btc_price", "")).strip(),
        "eth_usd": str(raw.get("eth_usd", "")).strip(),
        "eth_price": str(raw.get("eth_price", "")).strip(),
    }


def _normalize_crypto_refill(raw: dict) -> dict:
    rid = str(raw.get("id", "")).strip()
    return {
        "id": rid if rid else str(uuid.uuid4()),
        "date": str(raw.get("date", "")).strip(),
        "uah": str(raw.get("uah", "")).strip(),
        "usd": str(raw.get("usd", "")).strip(),
    }


def _normalize_crypto_rules(raw: dict | None) -> dict[str, str]:
    base = dict(_CRYPTO_DEFAULT_RULES)
    if isinstance(raw, dict):
        for k in base:
            if k in raw and raw[k] is not None:
                base[k] = str(raw[k]).strip()[:500]
    return base


def _sort_crypto_purchases(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: _parse_date(str(r.get("date", ""))))


def _sort_crypto_refills(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: _parse_date(str(r.get("date", ""))))


def _crypto_from_db_row(data: dict | None) -> InvestingCryptoResponse:
    d = data if isinstance(data, dict) else {}
    holdings_raw = d.get("holdings") if isinstance(d.get("holdings"), list) else []
    purchases_raw = d.get("purchase_rows") if isinstance(d.get("purchase_rows"), list) else []
    refills_raw = d.get("refill_rows") if isinstance(d.get("refill_rows"), list) else []
    rules_raw = d.get("rules")

    holdings = [_normalize_crypto_holding(x) for x in holdings_raw if isinstance(x, dict)]
    purchases = _sort_crypto_purchases([_normalize_crypto_purchase(x) for x in purchases_raw if isinstance(x, dict)])
    refills = _sort_crypto_refills([_normalize_crypto_refill(x) for x in refills_raw if isinstance(x, dict)])
    rules = _normalize_crypto_rules(rules_raw if isinstance(rules_raw, dict) else None)

    return InvestingCryptoResponse(
        holdings=[CryptoHolding(**h) for h in holdings],
        purchase_rows=[CryptoPurchaseRow(**p) for p in purchases],
        refill_rows=[CryptoRefillRow(**r) for r in refills],
        rules=CryptoRules(**rules),
        updated_at=None,
    )


@router.get("/crypto", response_model=InvestingCryptoResponse)
def get_crypto(
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    row = conn.execute(
        "SELECT data, updated_at FROM investing_crypto WHERE user_id = %s",
        (user.id,),
    ).fetchone()
    if not row:
        return _crypto_from_db_row({})
    data = row[0] if isinstance(row[0], dict) else {}
    payload = _crypto_from_db_row(data)
    return InvestingCryptoResponse(**(payload.model_dump() | {"updated_at": str(row[1])}))


@router.put("/crypto", response_model=InvestingCryptoResponse)
def update_crypto(
    payload: InvestingCryptoUpdateRequest,
    user: Annotated[UserResponse, Depends(get_current_user)],
    conn: Annotated[Connection, Depends(get_conn)],
):
    holdings = [_normalize_crypto_holding(h.model_dump()) for h in payload.holdings]
    purchases = _sort_crypto_purchases([_normalize_crypto_purchase(p.model_dump()) for p in payload.purchase_rows])
    refills = _sort_crypto_refills([_normalize_crypto_refill(r.model_dump()) for r in payload.refill_rows])
    rules = _normalize_crypto_rules(payload.rules.model_dump())

    data_obj = {
        "holdings": holdings,
        "purchase_rows": purchases,
        "refill_rows": refills,
        "rules": rules,
    }
    row = conn.execute(
        """
        INSERT INTO investing_crypto(user_id, data, updated_at)
        VALUES (%s, %s, NOW())
        ON CONFLICT (user_id)
        DO UPDATE SET data = EXCLUDED.data, updated_at = NOW()
        RETURNING updated_at
        """,
        (user.id, Json(data_obj)),
    ).fetchone()
    conn.commit()
    updated_at = str(row[0]) if row else None
    return InvestingCryptoResponse(
        holdings=[CryptoHolding(**h) for h in holdings],
        purchase_rows=[CryptoPurchaseRow(**p) for p in purchases],
        refill_rows=[CryptoRefillRow(**r) for r in refills],
        rules=CryptoRules(**rules),
        updated_at=updated_at,
    )

